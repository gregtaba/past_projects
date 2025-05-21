from docx import Document

# === Get User Input ===
employee_name = input("Enter the employee's full name: ")
employer_name = input("Enter the employer's name or firm: ")
effective_date = input("Enter the effective date (e.g. May 7, 2025): ")
duration = input("Enter the duration of the non-compete (e.g. 12 months): ")
restricted_area = input("Enter the geographic scope or region (e.g. State of New York): ")

# === Create Word Document ===
doc = Document()
doc.add_heading("Non-Compete Agreement", level=1)

doc.add_paragraph(f"This Non-Compete Agreement (\"Agreement\") is made effective as of {effective_date}, "
                  f"by and between {employer_name} (\"Employer\") and {employee_name} (\"Employee\").")

doc.add_paragraph("WHEREAS, the Employer desires to protect its legitimate business interests, and "
                  "the Employee has agreed to certain restrictions in consideration of continued employment;")

doc.add_paragraph("NOW, THEREFORE, in consideration of the mutual promises and covenants herein contained, "
                  "the parties agree as follows:")

doc.add_paragraph(f"1. **Non-Compete Obligation**\nThe Employee agrees that for a period of {duration} "
                  f"following the termination of their employment, they shall not engage, directly or indirectly, "
                  f"in any business that competes with the Employer's business within {restricted_area}.")

doc.add_paragraph("2. **Acknowledgment**\nThe Employee acknowledges that this restriction is reasonable in scope "
                  "and duration and is necessary to protect the legitimate business interests of the Employer.")

doc.add_paragraph("3. **Governing Law**\nThis Agreement shall be governed by and construed in accordance with "
                  "the laws of the applicable jurisdiction.")

doc.add_paragraph("IN WITNESS WHEREOF, the parties have executed this Agreement as of the date first written above.")

# === Signature Block ===
doc.add_paragraph("\n\n______________________________")
doc.add_paragraph(f"{employee_name} (Employee)")

doc.add_paragraph("\n\n______________________________")
doc.add_paragraph(f"{employer_name} (Employer)")

# === Save the Document ===
filename = f"Non_Compete_{employee_name.replace(' ', '_')}.docx"
doc.save(filename)

print(f"\nNon-compete agreement saved as: {filename}")
